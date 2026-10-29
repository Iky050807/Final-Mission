# ============================================================================
#  services/pptx_export.py
#  File .pptx SUNGGUHAN (python-pptx), 16:9, objek yang dapat diedit (kotak teks, bentuk) — BUKAN screenshot.
#  Seluruh posisi/ukuran berasal dari services/slide_layout.py, mesin yang sama dengan Preview dan
#  Mode Presentasi, sehingga tampilan di PowerPoint = tampilan di website.
# ============================================================================
import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN, MSO_AUTO_SIZE
from pptx.util import Pt

from services.slide_layout import W, H, LH, expand_deck

ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}


def _rgb(hex_color):
    h = (hex_color or "#000000").lstrip("#")
    if len(h) != 6:
        h = "000000"
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _draw_text(slide, el):
    box = slide.shapes.add_textbox(Pt(el["x"]), Pt(el["y"]), Pt(el["w"]), Pt(el["h"]))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR.get(el.get("valign", "top"), MSO_ANCHOR.TOP)
    first = True
    for p in el["paras"]:
        # satu elemen paragraf bisa memuat baris baru (kode): pisahkan menjadi paragraf PPTX agar baris tetap
        for line in str(p["text"]).split("\n"):
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.alignment = ALIGN.get(p.get("align", "left"), PP_ALIGN.LEFT)
            para.line_spacing = Pt(p["size"] * LH)           # jarak baris eksak = sama dgn preview
            para.space_after = Pt(p.get("after", 0)) if "\n" not in str(p["text"]) else Pt(0)
            run = para.add_run()
            run.text = line
            f = run.font
            f.size = Pt(p["size"])
            f.bold = bool(p.get("bold"))
            f.italic = bool(p.get("italic"))
            f.underline = bool(p.get("underline"))
            f.name = el.get("font", "Calibri")
            f.color.rgb = _rgb(p.get("color"))
    return box


def _draw_shape(slide, el):
    kind = MSO_SHAPE.OVAL if el["t"] == "ellipse" else (
        MSO_SHAPE.ROUNDED_RECTANGLE if el.get("radius", 0) > 4 else MSO_SHAPE.RECTANGLE)
    shp = slide.shapes.add_shape(kind, Pt(el["x"]), Pt(el["y"]), Pt(el["w"]), Pt(el["h"]))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = min(0.5, el["radius"] / max(1.0, min(el["w"], el["h"])))
    if el.get("fill"):
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(el["fill"])
    else:
        shp.fill.background()
    if el.get("line"):
        shp.line.color.rgb = _rgb(el["line"])
        shp.line.width = Pt(el.get("lw", 1.25))
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def build_pptx(config, members=None, group_name=""):
    deck = expand_deck(config, members or [""] * 10, group_name)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(W), Pt(H)          # 960 x 540 pt = 13.333 x 7.5 in (16:9)
    blank = prs.slide_layouts[6]
    for page in deck["pages"]:
        slide = prs.slides.add_slide(blank)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(page["bg"])
        for el in page["elements"]:
            if el["t"] == "text":
                _draw_text(slide, el)
            else:
                _draw_shape(slide, el)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
